"""通用 Agent 能力：附件、视觉、导出/写文件、定时任务、改问题重答、对话搜索、对话管理（归档/删除撤销/自动标题）、通用问答提示词。全部离线。"""
import io
import json
import threading
import time
import urllib.error
import urllib.request
import zipfile
from datetime import datetime
from http.server import ThreadingHTTPServer

import pytest


@pytest.fixture
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("STOCK_AGENT_HOME", str(tmp_path))
    for k in ("LLM_VISION", "WEB_TOKEN"):
        monkeypatch.delenv(k, raising=False)
    return tmp_path


def _pdf(text="Revenue 123.45 Stoa"):
    """手写一个最小的带文字 PDF。"""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offs.append(out.tell())
        out.write(b"%d 0 obj\n" % i + o + b"\nendobj\n")
    x = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs))
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x))
    return out.getvalue()


def _docx(paras):
    body = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paras)
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document xmlns:w="x"><w:body>{body}</w:body></w:document>')
    return b.getvalue()


PNG = bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082")


# ---------------- 1. 附件 ----------------
def test_uploads_extract_each_kind(home):
    from stock_agent import uploads
    c = uploads.save("持仓.csv", "代码,数量,成本\n600519,100,1500.5\n0700.HK,200,320\n".encode("utf-8"))
    assert c["kind"] == "table" and c["rows"] == 2
    meta = uploads.get(c["id"])
    assert "1500.5" in meta["text"] and "数值列统计" in meta["text"]
    import pandas as pd
    b = io.BytesIO()
    with pd.ExcelWriter(b) as w:
        pd.DataFrame({"a": [1, 2, 3]}).to_excel(w, sheet_name="一", index=False)
        pd.DataFrame({"b": ["x"]}).to_excel(w, sheet_name="二", index=False)
    x = uploads.save("book.xlsx", b.getvalue())
    assert x["sheets"] == ["一", "二"] and x["rows"] == 4
    d = uploads.save("纪要.docx", _docx(["第一段 营收 88 亿", "第二段"]))
    assert "营收 88 亿" in uploads.get(d["id"])["text"]
    p = uploads.save("r.pdf", _pdf())
    assert p["kind"] == "pdf" and p["pages"] == 1 and "123.45" in uploads.get(p["id"])["text"]
    t = uploads.save("note.md", "# 标题\n内容".encode("gb18030"))
    assert "内容" in uploads.get(t["id"])["text"]
    im = uploads.save("k.png", PNG)
    assert im["kind"] == "image" and uploads.file_path(im["id"]).exists()
    assert "error" in uploads.save("evil.exe", b"MZ")
    assert "error" in uploads.save("a.txt", b"")
    assert uploads.get("../../etc") is None and uploads.file_path("zzz") is None


def test_upload_size_limit(home, monkeypatch):
    from stock_agent import uploads
    monkeypatch.setattr(uploads, "MAX_BYTES", 10)
    assert "太大" in uploads.save("a.txt", b"x" * 11)["error"]


def test_marker_roundtrip_and_for_model(home):
    from stock_agent import uploads
    a = uploads.save("a.csv", b"x,y\n1,2\n")["id"]
    im = uploads.save("k.png", PNG)["id"]
    text = uploads.mark("看看这个", [a, im, "bad!"])
    body, ids = uploads.split(text)
    assert body == "看看这个" and ids == [a, im]
    content, images, corpus = uploads.for_model(text, vision=True)
    assert "<<<附件「a.csv」" in content and "run_python" in content and images[0]["url"].startswith("data:image/png;base64,")
    assert "x,y" in corpus
    content2, images2, _ = uploads.for_model(text, vision=False)
    assert not images2 and "不支持看图" in content2
    assert uploads.for_model("普通问题")[0] == "普通问题"


def test_vision_heuristic(monkeypatch):
    from stock_agent import uploads
    assert uploads.vision_capable("gpt-4o-mini") and uploads.vision_capable("claude-sonnet-4") and uploads.vision_capable("qwen-vl-max")
    assert not uploads.vision_capable("deepseek-chat") and not uploads.vision_capable("mock")
    assert uploads.vision_capable("deepseek-chat", {"vision": True})
    monkeypatch.setenv("LLM_VISION", "0")
    assert not uploads.vision_capable("gpt-4o")


def test_provider_parts_convert_images():
    from stock_agent import providers
    c = [{"type": "text", "text": "hi"}, {"type": "image_url", "image_url": {"url": "data:image/png;base64,QUJD"}}]
    assert providers._parts(c) == [("text", "hi"), ("image", "image/png", "QUJD")]
    assert providers._parts("plain") == [("text", "plain")]


class _Spy:
    """记录发给模型的消息；第一次见到图片时可选择报错（模拟不支持视觉的模型）。"""
    ready = True

    def __init__(self, reject_images=False, reply="好的，已读。"):
        self.calls, self.reject, self.reply = [], reject_images, reply
        self.chain = [type("M", (), {"model": "gpt-4o"})()]

    context = 128000

    def chat(self, messages, tools=None, temperature=0.3, timeout=120):
        self.calls.append(json.loads(json.dumps(messages)))
        if self.reject and any(isinstance(m.get("content"), list) for m in messages):
            raise RuntimeError("400 image input not supported")
        return {"role": "assistant", "content": self.reply, "tool_calls": None}


def _agent(home, llm, session=None):
    from stock_agent.agent import Agent
    from stock_agent.memory import Memory
    ag = Agent(Memory(home), channel="test", session=session)
    ag.llm = llm
    return ag


def test_agent_sends_attachment_and_image(home):
    from stock_agent import uploads
    a = uploads.save("a.csv", b"x,y\n1,2\n")["id"]
    im = uploads.save("k.png", PNG)["id"]
    spy = _Spy()
    ag = _agent(home, spy)
    ag.ask(uploads.mark("总结附件", [a, im]))
    last_user = [m for m in spy.calls[-1] if m["role"] == "user"][-1]
    assert isinstance(last_user["content"], list) and last_user["content"][1]["type"] == "image_url"
    assert "<<<附件「a.csv」" in last_user["content"][0]["text"]
    assert not any("_images" in m for m in ag.messages)  # 答完就不再带图
    # 历史里只存标记
    assert "[[files:" in ag.mem.history(ag.session)[0]["text"]
    # 恢复旧对话：附件文字展开回来
    ag2 = _agent(home, _Spy(), session=ag.session)
    assert any("<<<附件「a.csv」" in (m.get("content") or "") for m in ag2.messages if m["role"] == "user")


def test_agent_image_fallback_when_model_rejects(home):
    from stock_agent import uploads
    im = uploads.save("k.png", PNG)["id"]
    spy = _Spy(reject_images=True)
    ag = _agent(home, spy)
    ag.ask(uploads.mark("这张图是什么", [im]))
    last = [m for m in spy.calls[-1] if m["role"] == "user"][-1]
    assert isinstance(last["content"], str) and "支持视觉" in last["content"]
    assert ag._vision_ok is False


# ---------------- 3. 导出 ----------------
def test_write_xlsx_csv_md_and_resolve(home):
    from stock_agent import exports
    r = exports.write("估值表", "xlsx", sheets=[{"name": "A股", "columns": ["代码", "PE", "备注"], "rows": [["600519", 22.5, "=HYPERLINK(1)"]]},
                                                {"name": "港股", "rows": [{"代码": "0700.HK", "PE": 18}]}])
    assert r["ok"] and r["url"].startswith("/api/exports/") and r["markdown_link"].startswith("[")
    from openpyxl import load_workbook
    wb = load_workbook(exports.resolve(r["file"]))
    assert wb.sheetnames == ["A股", "港股"] and wb["A股"]["B2"].value == 22.5
    assert wb["A股"]["C2"].value.startswith("'=")  # 防公式注入
    assert wb["港股"]["A1"].value == "代码"
    c = exports.write("x.csv", None, columns=["a"], rows=[[1], [2]])
    assert exports.resolve(c["file"]).read_bytes().startswith(b"\xef\xbb\xbf")
    m = exports.write("../../etc/passwd", "md", content="# hi")
    assert "/" not in m["file"] and exports.resolve(m["file"]).read_text(encoding="utf-8") == "# hi"
    assert exports.write("y", "exe")["error"]
    for bad in ("../x.md", "..", "/etc/passwd", "%2e%2e%2fx", ".hidden", "a/b.md"):
        assert exports.resolve(bad) is None, bad
    assert exports.listing()


def test_write_file_tool_and_conversation_md(home):
    ag = _agent(home, _Spy())
    res = ag.run_tool("write_file", {"filename": "表", "format": "csv", "columns": ["a"], "rows": [[1]]})
    assert res["ok"] and res["file"].endswith(".csv")
    from stock_agent import exports, uploads
    a = uploads.save("a.csv", b"x\n1\n")["id"]
    md = exports.conversation_md("测试", [{"role": "user", "text": uploads.mark("问题一", [a]), "time": "2026-10-31 10:00"},
                                          {"role": "assistant", "text": "回答 [1]\n[[cards:abc]]"}])
    assert md.startswith("# 测试") and "## 1. 问题一" in md and "附件：a.csv" in md and "[[cards" not in md


# ---------------- 4. 定时任务 ----------------
def _t(s):
    from stock_agent.tasks import CN
    return datetime.strptime(s, "%Y-%m-%d %H:%M").replace(tzinfo=CN)


def test_task_schedules(home):
    from stock_agent import tasks
    sat = _t("2026-10-31 10:00")  # 周六
    assert tasks.next_run({"kind": "daily", "time": "09:00"}, sat) == _t("2026-11-01 09:00")
    assert tasks.next_run({"kind": "weekdays", "time": "09:00"}, sat) == _t("2026-11-02 09:00")
    assert tasks.next_run({"kind": "weekly", "time": "20:00", "weekdays": [5]}, sat) == _t("2026-10-31 20:00")
    assert tasks.next_run({"kind": "interval", "every_minutes": 90}, sat) == _t("2026-10-31 11:30")
    assert tasks.next_run({"kind": "once", "at": "2026-10-31 09:00"}, sat) is None
    assert tasks.add("x", "daily", "25:00")["error"] and tasks.add("x", "interval", every_minutes=5)["error"]
    assert tasks.add("", "daily", "9:00")["error"] and tasks.add("x", "weekly", "9:00")["error"]
    assert tasks.add("x", "once", at="2001-01-01 09:00")["error"]
    ok = tasks.add("总结自选股", "daily", "9点", title="早间")
    assert ok["ok"] and ok["schedule"] == "每天 09:00"


def test_task_tick_runs_and_notifies(home):
    from stock_agent import alerts, tasks
    r = tasks.add("搜一下今天的 AI 新闻", "daily", "09:00")
    seen, pushed = [], []
    t0 = tasks._parse(tasks.load()[0]["next_run"])
    n = tasks.tick(lambda sid, text: seen.append((sid, text)) or "今天 AI 新闻三条", push=lambda *a: pushed.append(a), at=t0)
    assert n == 1 and seen[0][0] == "task-" + r["id"] and seen[0][1].startswith(tasks.TASK_PREFIX) and "AI 新闻" in seen[0][1]
    t = tasks.load()[0]
    assert t["runs"] == 1 and t["last_status"] == "ok" and tasks._parse(t["next_run"]) > t0
    ev = alerts.load()["events"][-1]
    assert ev["kind"] == "task" and ev["session"] == "task-" + r["id"] and pushed
    assert tasks.tick(lambda *a: "x", push=lambda *a: None, at=t0) == 0  # 不重复触发
    # 停机太久：跳过不补
    late = tasks._parse(t["next_run"]) + tasks.CATCH_UP + tasks.timedelta(minutes=5)
    assert tasks.tick(lambda *a: 1 / 0, push=lambda *a: None, at=late) == 0
    assert tasks.remove(r["id"])["ok"] and tasks.load() == []


def test_task_once_deactivates_and_agent_tools(home):
    from stock_agent import tasks
    ag = _agent(home, _Spy())
    res = ag.run_tool("task_create", {"prompt": "提醒我交房租", "kind": "once", "at": "2099-01-01 10:00", "title": "房租"})
    assert res["ok"]
    lst = ag.run_tool("task_list", {})
    assert lst["tasks"][0]["title"] == "房租"
    tasks.tick(lambda *a: "已提醒", push=lambda *a: None, at=_t("2099-01-01 10:00"))
    assert tasks.load()[0]["active"] is False
    ag._cur_text = tasks.TASK_PREFIX + "「x」]\n..."
    assert "error" in ag.run_tool("task_create", {"prompt": "再建一个", "kind": "daily", "time": "09:00"})
    assert ag.run_tool("task_delete", {"id": res["id"]})["ok"]


# ---------------- 2/6. 改问题重答、搜索 ----------------
def test_truncate_turn_and_search_snippet(home):
    from stock_agent.memory import Memory
    m = Memory(home)
    s = "web-abc"
    for q, a in (("第一问 茅台", "答一"), ("第二问 腾讯控股的回购", "答二"), ("第三问", "答三")):
        m.log(s, "user", q); m.log(s, "assistant", a)
    m.save_digest(s, 10 ** 9, "纪要", [], 2)
    assert m.truncate_turn(s, 9) is None
    cut = m.truncate_turn(s, 2)
    assert cut["text"] == "第二问 腾讯控股的回购"
    assert [h["text"] for h in m.history(s)] == ["第一问 茅台", "答一"]
    assert m.digest(s) is None
    assert m.sessions(query="腾讯控股") == []  # 全文索引也删了
    m.log(s, "user", "再问：腾讯控股今年回购了多少")
    hit = m.sessions(query="腾讯控股")[0]
    assert hit["id"] == "abc" and "腾讯控股" in hit["hit"] and hit["hit_id"]


# ---------------- 对话管理 ----------------
def test_archive_softdelete_undo_and_titles(home, monkeypatch):
    from stock_agent import memory as M, titles
    m = M.Memory(home)
    for sid in ("a", "b"):
        m.log("web-" + sid, "user", "问 " + sid); m.log("web-" + sid, "assistant", "答")
    m.rename_session("web-a", pinned=True)
    m.rename_session("web-a", archived=True)
    assert [x["id"] for x in m.sessions()] == ["b"]
    arch = m.sessions(archived=True)
    assert arch[0]["id"] == "a" and arch[0]["archived"] and not arch[0]["pinned"]
    m.rename_session("web-a", archived=False)
    assert {x["id"] for x in m.sessions()} == {"a", "b"}
    assert m.soft_delete("web-b") and [x["id"] for x in m.sessions()] == ["a"]
    assert m.restore_session("web-b") and len(m.sessions()) == 2
    m.soft_delete("web-b")
    monkeypatch.setattr(M, "UNDO_SECONDS", -1)
    m.purge_deleted()
    assert not m.history("web-b") and not m.restore_session("web-b")
    assert not m.soft_delete("web-nope")
    # 自动标题
    assert titles.heuristic("请问 腾讯控股最近为什么一直跌？") == "腾讯控股最近为什么一直跌"
    assert titles.heuristic("How do I reverse a linked list in Python please help me") .startswith("How do I reverse")
    assert titles.clean("标题：《Python 链表反转》") == "Python 链表反转"

    class T:
        def chat(self, *a, **k):
            return {"content": "「茅台估值复盘」"}
    monkeypatch.delenv("STOA_MOCK_LLM", raising=False)
    assert titles.auto_title(m, "web-a", "茅台贵不贵", "……", T()) == "茅台估值复盘"
    assert m.sessions()[0]["auto_title"] is True
    assert titles.auto_title(m, "web-a", "x", "y", T()) is None  # 已有标题不覆盖
    m.rename_session("web-a", title="我的名字")
    assert m.sessions()[0]["title"] == "我的名字" and not m.sessions()[0]["auto_title"]


# ---------------- 5. 通用问题 ----------------
def test_system_prompt_general_and_disclaimer_scope():
    from stock_agent import agent as A
    assert "通用问题" in A.SYSTEM and "不要拒绝" in A.SYSTEM and "write_file" in A.SYSTEM and "task_create" in A.SYSTEM
    assert not A.wants_disclaimer("推荐一本学 Python 的书")
    assert not A.wants_disclaimer("帮我翻译：the stock of knowledge")  # 英文 stock 不算
    assert A.wants_disclaimer("推荐一只股票") and A.wants_disclaimer("借钱加杠杆炒腾讯可以吗") and A.wants_disclaimer("600519 能买吗")
    names = {t["name"] for t in A.TOOLS}
    assert {"write_file", "task_create", "task_list", "task_delete"} <= names


def test_general_question_no_disclaimer(home):
    spy = _Spy(reply="def rev(s):\n    return s[::-1]")
    ag = _agent(home, spy)
    out = ag.ask("用 Python 写一个反转字符串的函数，推荐写法")
    assert "不构成投资建议" not in out


# ---------------- HTTP 端到端 ----------------
@pytest.fixture
def server(home, monkeypatch):
    monkeypatch.setenv("STOA_MOCK_LLM", "1")
    from stock_agent import serve
    from stock_agent.memory import Memory
    monkeypatch.setattr(serve, "MEM", Memory(home))
    serve.SESSIONS.clear()
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), serve.Web)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}", serve
    httpd.shutdown()
    serve.SESSIONS.clear()


_OPEN = urllib.request.build_opener(urllib.request.ProxyHandler({})).open  # 沙箱有 http_proxy，本机请求不走代理


def _req(url, body=None, raw=None, method=None):
    data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
    r = urllib.request.Request(url, data=data, method=method or ("POST" if data is not None else "GET"))
    try:
        with _OPEN(r, timeout=30) as f:
            return f.status, f.read(), dict(f.headers)
    except urllib.error.HTTPError as e:
        return e.code, e.read(), dict(e.headers)


def test_http_upload_chat_edit_export_tasks(server):
    base, serve = server
    st, b, _ = _req(base + "/api/upload?name=" + urllib.request.quote("持仓.csv"), raw="代码,成本\n600519,1500\n".encode())
    up = json.loads(b)
    assert st == 200 and up["kind"] == "table"
    st, b, h = _req(base + "/api/uploads/" + up["id"])
    assert st == 200 and h["Content-Disposition"].startswith("attachment")
    assert _req(base + "/api/upload?name=a.exe", raw=b"MZ")[0] == 400
    st, b, _ = _req(base + "/api/chat", {"text": "第一问 你好", "session": "s1", "files": [up["id"]]})
    assert st == 200 and json.loads(b)["reply"]
    _req(base + "/api/chat", {"text": "第二问", "session": "s1"})
    hist = serve.MEM.history("web-s1")
    assert sum(h["role"] == "user" for h in hist) == 2 and "[[files:" in hist[0]["text"]
    # 改第 2 问
    st, b, _ = _req(base + "/api/chat", {"text": "第二问改过", "session": "s1", "edit": 2})
    assert st == 200
    qs = [h["text"] for h in serve.MEM.history("web-s1") if h["role"] == "user"]
    assert len(qs) == 2 and qs[1] == "第二问改过"
    assert _req(base + "/api/chat", {"text": "x", "session": "s1", "edit": 9})[0] == 400
    # 自动标题（mock 模式按规则）
    for _ in range(50):
        if serve.MEM.has_title("web-s1"):
            break
        time.sleep(.05)
    j = json.loads(_req(base + "/api/sessions")[1])
    assert j["sessions"][0]["title"] == "第一问 你好" and j["sessions"][0]["auto_title"]
    # 导出对话
    st, b, h = _req(base + "/api/sessions/s1/export.md")
    assert st == 200 and b.decode().startswith("# ") and "第二问改过" in b.decode() and "attachment" in h["Content-Disposition"]
    # 导出文件下载 + 路径限制
    from stock_agent import exports
    f = exports.write("t", "md", content="ok")["file"]
    assert _req(base + "/api/exports/" + urllib.request.quote(f))[1] == b"ok"
    assert _req(base + "/api/exports/..%2F..%2Fdata%2Fsessions.db")[0] == 404
    assert _req(base + "/api/exports/%2Fetc%2Fpasswd")[0] == 404
    assert json.loads(_req(base + "/api/exports")[1])["files"]
    # 定时任务接口
    from stock_agent import tasks
    tid = tasks.add("说一句早安", "daily", "09:00")["id"]
    assert json.loads(_req(base + "/api/tasks")[1])["tasks"][0]["id"] == tid
    assert json.loads(_req(base + "/api/tasks/toggle", {"id": tid, "active": False})[1])["ok"]
    assert tasks.load()[0]["active"] is False
    assert json.loads(_req(base + "/api/tasks/delete", {"id": tid})[1])["ok"]


def test_http_archive_delete_undo(server):
    base, serve = server
    serve.MEM.log("web-z", "user", "问"); serve.MEM.log("web-z", "assistant", "答")
    _req(base + "/api/sessions/rename", {"id": "z", "archived": True})
    j = json.loads(_req(base + "/api/sessions")[1])
    assert j["sessions"] == [] and j["archived_count"] == 1
    assert json.loads(_req(base + "/api/sessions?archived=1")[1])["sessions"][0]["id"] == "z"
    d = json.loads(_req(base + "/api/sessions/delete", {"id": "z"})[1])
    assert d["undo"] > 0 and json.loads(_req(base + "/api/sessions?archived=all")[1])["sessions"] == []
    assert json.loads(_req(base + "/api/sessions/restore", {"id": "z"})[1])["ok"]
    assert json.loads(_req(base + "/api/sessions?archived=1")[1])["sessions"]
    _req(base + "/api/sessions/delete", {"id": "z"})
    _req(base + "/api/sessions/purge", {"id": "z"})
    assert not serve.MEM.history("web-z")


def test_web_ui_wired():
    html = open("stock_agent/web.html", encoding="utf-8").read()
    for s in ('id="att"', "function uploadFiles", "dragover", "'paste'", "function editTurn", "function copyAnswer", "export.md",
              "function loadTasks", "/api/tasks", "?archived=", "/api/sessions/restore", "data-act=\"arc\"", "hit_id"):
        assert s in html, s
    en = json.load(open("stock_agent/assets/i18n-en.json", encoding="utf-8"))
    for k in ("归档", "撤销", "定时任务", "附件", "复制", "重新生成", "导出 Markdown"):
        assert k in en, k
