"""导出：Agent 写文件（Excel / CSV / Markdown / 文本 / JSON）到 数据目录/exports，网页给下载链接；
以及把一段对话导出成 Markdown。下载接口只认 exports 目录下的文件名，不接受路径。"""
import csv
import io
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from .memory import home

CN = timezone(timedelta(hours=8))
FORMATS = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8",
           "md": "text/markdown; charset=utf-8", "txt": "text/plain; charset=utf-8", "json": "application/json; charset=utf-8"}
MAX_ROWS = 100000
MAX_CHARS = 2_000_000
NAME_RE = re.compile(r"^[^/\\\x00-\x1f]{1,160}$")


def root() -> Path:
    p = home() / "exports"
    p.mkdir(parents=True, exist_ok=True)
    return p


def _clean_stem(name: str) -> str:
    stem = Path(str(name or "").replace("\\", "/")).name
    stem = re.sub(r"\.(xlsx|csv|md|txt|json)$", "", stem, flags=re.I)
    stem = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", stem).strip(" ._")
    return stem[:80] or "stoa"


def _unique(stem: str, ext: str) -> Path:
    day = datetime.now(CN).strftime("%Y%m%d")
    p = root() / f"{stem}-{day}.{ext}"
    i = 2
    while p.exists():
        p = root() / f"{stem}-{day}-{i}.{ext}"
        i += 1
    return p


def _rows(spec: dict) -> tuple[list, list]:
    cols = [str(c) for c in (spec.get("columns") or [])]
    rows = spec.get("rows") or []
    if rows and isinstance(rows[0], dict):  # [{"列": 值}] 也认
        if not cols:
            for r in rows:
                for k in r:
                    if str(k) not in cols:
                        cols.append(str(k))
        rows = [[{str(k): v for k, v in r.items()}.get(c) for c in cols] for r in rows]
    rows = [list(r) if isinstance(r, (list, tuple)) else [r] for r in rows][:MAX_ROWS]
    return cols, rows


def _cell(v):
    if v is None or isinstance(v, (int, float, bool)):
        return v
    s = str(v)
    if s[:1] in ("=", "+", "-", "@") and not re.fullmatch(r"[-+]?\d+(\.\d+)?%?", s):  # 防公式注入
        s = "'" + s
    return s


def write(filename: str, fmt: str | None = None, content: str | None = None, columns=None, rows=None, sheets=None,
          replace: bool = False, session: str | None = None, _approved: bool = False) -> dict:
    """写一个文件。fmt 不给就看扩展名。xlsx 用 sheets=[{name, columns, rows}] 或 columns+rows；csv 用 columns+rows；md/txt/json 用 content。
    默认每次写新文件（文件名带日期，重名加序号，从不覆盖）。replace=True 写到固定文件名「名字.扩展名」：
    文件已存在时不直接覆盖，而是挂一条待确认（重启后仍在），用户批准后才写。"""
    ext = (fmt or (str(filename).rsplit(".", 1)[-1] if "." in str(filename) else "md")).lower().lstrip(".")
    if ext == "markdown":
        ext = "md"
    if ext not in FORMATS:
        return {"error": f"只支持 {', '.join(FORMATS)}"}
    if replace:
        p = root() / f"{_clean_stem(filename)}.{ext}"
        if p.exists() and not _approved:
            from . import ledger
            a = ledger.approval_add("file_overwrite", f"覆盖已有文件 {p.name}",
                                    {"filename": filename, "fmt": ext, "content": content, "columns": columns, "rows": rows, "sheets": sheets,
                                     "file": p.name}, session=session, dedupe=f"file_overwrite:{p.name}")
            return {"pending_approval": a["id"], "file": p.name,
                    "note": f"{p.name} 已存在，覆盖要用户确认：已挂起（编号 {a['id']}），用户在「待确认」里批准或在对话里明确说确认后才会写入。告诉用户这一点，不要说已经写好。"}
    else:
        p = _unique(_clean_stem(filename), ext)
    if ext == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            return {"error": "写 Excel 需要 openpyxl（pip install openpyxl）；可以改成 csv"}
        sh = sheets or [{"name": "Sheet1", "columns": columns or [], "rows": rows or []}]
        wb = Workbook()
        wb.remove(wb.active)
        n = 0
        for i, s in enumerate(sh[:20]):
            title = re.sub(r"[\[\]:*?/\\]", "_", str(s.get("name") or f"Sheet{i + 1}"))[:31] or f"Sheet{i + 1}"
            ws = wb.create_sheet(title)
            cols, rs = _rows(s)
            if cols:
                ws.append(cols)
                for c in ws[1]:
                    c.font = Font(bold=True)
                ws.freeze_panes = "A2"
            for r in rs:
                ws.append([_cell(v) for v in r])
            n += len(rs)
            for j, c in enumerate(cols[:50], 1):  # 列宽按表头和前 50 行估
                w = max([len(str(c))] + [len(str(r[j - 1])) for r in rs[:50] if len(r) >= j])
                ws.column_dimensions[ws.cell(1, j).column_letter].width = min(48, max(8, w * 1.3))
        wb.save(p)
        info = f"{len(sh)} 个工作表、{n} 行"
    elif ext == "csv":
        cols, rs = _rows({"columns": columns, "rows": rows})
        if not cols and not rs and content:
            p.write_text(str(content)[:MAX_CHARS], encoding="utf-8-sig")
            info = "CSV"
        else:
            buf = io.StringIO()
            w = csv.writer(buf)
            if cols:
                w.writerow(cols)
            w.writerows([[_cell(v) for v in r] for r in rs])
            p.write_text(buf.getvalue(), encoding="utf-8-sig")  # 带 BOM，Excel 直接打开中文不乱码
            info = f"{len(rs)} 行"
    else:
        body = content if isinstance(content, str) else json.dumps(content if content is not None else {"columns": columns, "rows": rows},
                                                                 ensure_ascii=False, indent=2)
        p.write_text(body[:MAX_CHARS], encoding="utf-8")
        info = f"{len(body)} 字"
    url = "/api/exports/" + p.name
    return {"ok": True, "file": p.name, "url": url, "size": p.stat().st_size, "info": info,
            "markdown_link": f"[{p.name}]({url})",
            "note": "在回答里原样写出 markdown_link，网页会显示成下载按钮。不要编造别的链接。"}


def resolve(name: str) -> Path | None:
    """下载：只接受 exports 目录里的文件名（不含路径）。"""
    from urllib.parse import unquote
    name = unquote(str(name or ""))
    if not NAME_RE.match(name) or name in (".", "..") or name.startswith("."):
        return None
    base = root().resolve()
    p = (base / name).resolve()
    if p.parent != base or not p.is_file():
        return None
    return p


def ctype(p: Path) -> str:
    return FORMATS.get(p.suffix.lstrip(".").lower(), "application/octet-stream")


def listing(limit: int = 50) -> list[dict]:
    fs = sorted((f for f in root().iterdir() if f.is_file()), key=lambda f: f.stat().st_mtime, reverse=True)[:limit]
    return [{"file": f.name, "url": "/api/exports/" + f.name, "size": f.stat().st_size,
             "time": datetime.fromtimestamp(f.stat().st_mtime, CN).strftime("%Y-%m-%d %H:%M")} for f in fs]


def conversation_md(title: str, messages: list[dict]) -> str:
    """一段对话 → Markdown：问题用二级标题，回答原样（保留引用和来源列表），去掉网页内部标记。"""
    from . import uploads
    out = [f"# {title or 'Stoa 对话'}", "", f"> 导出于 {datetime.now(CN):%Y-%m-%d %H:%M}（北京时间） · Stoa", ""]
    n = 0
    for m in messages:
        if m.get("role") == "user":
            n += 1
            body, ids = uploads.split(m.get("text") or "")
            names = [x["name"] for x in (uploads.get(i) for i in ids) if x]
            out += [f"## {n}. {body.strip().splitlines()[0][:80] if body.strip() else '（附件）'}", ""]
            rest = "\n".join(body.strip().splitlines()[1:]).strip()
            if rest:
                out += [rest, ""]
            if names:
                out += ["附件：" + "、".join(names), ""]
            out += [f"*{m.get('time', '')}*", ""]
        elif m.get("role") == "assistant":
            out += [clean_answer(m.get("text") or ""), "", "---", ""]
    return "\n".join(out).rstrip() + "\n"


def clean_answer(text: str) -> str:
    return re.sub(r"\n?\[\[(cards|debate|files):[^\]]*\]\]", "", text or "").strip()


def _approved_overwrite(pl: dict) -> dict:
    return write(pl.get("filename") or pl.get("file") or "stoa", pl.get("fmt"), pl.get("content"), pl.get("columns"), pl.get("rows"),
                 pl.get("sheets"), replace=True, _approved=True)


def _register():
    from . import ledger
    ledger.HANDLERS["file_overwrite"] = _approved_overwrite


_register()
