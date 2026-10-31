"""投研 PPT（v0.37，P3）：write_file(format="pptx", deck={...}) → 固定 8 页 16:9 投研模板，python-pptx 生成，原生图表（数据在图表里，可在
WPS / PowerPoint 里直接编辑）。

1 结论（看多/看空/震荡 + 置信度 + 三条依据）  2 价格与估值（指标 + 价格走势）  3 财务（三年关键项 + 增速）
4 利润结构瀑布  5 Agora 辩论摘要  6 风险清单  7 同行对比（sweep_filings 矩阵）  8 失效条件与复盘日期

缺的部分：给了 symbol 时先从本机的 Agora 辩论纪要（data/debates）与判断账本补结论、辩论、风险、失效条件与复盘日期；仍缺的页照样出，
写「本页数据未提供」，不编数字。生成后用 python-pptx 读回校验：8 页齐、页标题对、输入里的每个数字都出现在页面文字或图表数据里。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

TITLES = ["结论", "价格与估值", "财务", "利润结构", "多空辩论摘要", "风险清单", "同行对比", "失效条件与复盘日期"]
LATIN = ["Conclusio", "Pretium", "Rationes", "Structura", "Agora", "Pericula", "Comparatio", "Revisio"]
INK, GOLD, TERRA, OLIVE, MUTE, PAPER, LINE = "181715", "9E7B56", "B3543A", "56684F", "6F6A60", "F5F2EB", "D9D2C3"
SERIF, SANS = "Georgia", "Microsoft YaHei"
EMPTY = "本页数据未提供（Stoa 不替你补数字）"
VIEW_COLOR = {"看多": TERRA, "看涨": TERRA, "看空": OLIVE, "看跌": OLIVE}


def available() -> bool:
    try:
        import pptx  # noqa: F401
        return True
    except ImportError:
        return False


# ---------------------------------------------------------------- 从本机记录补齐
def autofill(deck: dict) -> dict:
    """deck.symbol 有、某些页没给 → 用最近一次该标的的 Agora 辩论纪要与判断账本补；补进来的页记在 filled 里（页脚注明来源）。"""
    d = dict(deck or {})
    sym = str(d.get("symbol") or "").strip()
    filled = []
    if not sym:
        return {**d, "_filled": filled}
    rec = None
    try:
        from .debate import debates_dir
        cands = []
        for p in debates_dir().glob("*.json"):
            try:
                r = json.loads(p.read_text(encoding="utf-8"))
            except Exception:
                continue
            if str(r.get("symbol") or "").upper() == sym.upper() and r.get("mode", "live") == "live":
                cands.append((r.get("created") or "", r))
        if cands:
            rec = max(cands, key=lambda x: x[0])[1]
    except Exception:
        rec = None
    if rec:
        v = rec.get("verdict") or {}
        src = f"Agora 辩论纪要 {rec.get('id')}（{rec.get('created')}）"
        if not d.get("conclusion") and v.get("view"):
            d["conclusion"] = {"view": v.get("view"), "confidence": v.get("confidence"), "horizon_days": v.get("horizon_days"),
                               "reasons": [re.sub(r"\[\d+\]", "", str(x)) for x in (v.get("key_points") or [])][:3], "source": src}
            filled.append("结论")
        if not d.get("debate") and rec.get("debate"):
            last = rec["debate"][-1]
            pts = lambda side: [re.sub(r"\[\d+\]", "", str(p)) for p in ((last.get(side) or {}).get("points") or
                                ([(last.get(side) or {}).get("argument")] if (last.get(side) or {}).get("argument") else []))][:3]
            risk = rec.get("risk") or {}
            d["debate"] = {"bull": pts("bull"), "bear": pts("bear"), "verdict": re.sub(r"\[\d+\]", "", str(v.get("summary") or "")),
                           "risk": f"风控：{risk.get('decision') or '—'}；单票仓位上限 {risk.get('max_position_pct', '—')}%", "source": src}
            filled.append("多空辩论摘要")
        if not d.get("risks") and rec.get("risks"):
            d["risks"] = [{"level": r.get("level"), "text": re.sub(r"\[\d+\]", "", str(r.get("text")))} for r in rec["risks"]][:8]
            filled.append("风险清单")
        if not d.get("invalidation") and v.get("invalidation"):
            d["invalidation"] = {"conditions": [re.sub(r"\[\d+\]", "", str(v["invalidation"]))], "source": src}
            filled.append("失效条件")
    inv = d.get("invalidation") or {}
    if not inv.get("review_date"):
        try:
            from .memory import Memory
            calls = [c for c in Memory().recent_calls(200) if str(c.get("symbol") or "").upper() == sym.upper() and c.get("status") == "open"]
            if calls:
                c = calls[0]
                d["invalidation"] = {**inv, "review_date": c.get("due"), "call": f"{c.get('view')} · 记账于 {str(c.get('ts'))[:10]} · 记账价 {c.get('price')}",
                                     "conditions": inv.get("conditions") or [], "source": inv.get("source") or "Stoa 判断账本"}
                filled.append("复盘日期")
        except Exception:
            pass
    return {**d, "_filled": filled}


# ---------------------------------------------------------------- 绘制
def _num_text(v) -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        if abs(v) >= 1e8:
            return f"{v / 1e8:,.2f} 亿"
        if float(v).is_integer():
            return str(int(v)) if 1900 <= v <= 2100 else f"{int(v):,}"  # 年份不加千分位
        return f"{v:,.2f}" if abs(v) >= 100 else f"{v:.4g}" if abs(v) < 1 else f"{v:.2f}"
    return str(v)


class _Deck:
    def __init__(self):
        from pptx import Presentation
        from pptx.util import Emu
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
        self.W, self.H = self.prs.slide_width, self.prs.slide_height

    def rgb(self, h):
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(h)

    def text(self, sl, x, y, w, h, s, size=14, color=INK, bold=False, font=SANS, align=None, italic=False):
        from pptx.util import Pt
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        lines = s if isinstance(s, list) else [s]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = str(ln)
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
            f.color.rgb = self.rgb(color)
            rpr = r._r.get_or_add_rPr()  # 中文字体（east asian）
            ea = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {"typeface": SANS if font == SANS else "SimSun"})
            rpr.append(ea)
            if align is not None:
                p.alignment = align
            p.space_after = Pt(size * .45)
        return tb

    def rule(self, sl, x, y, w, color=GOLD, h=None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu
        s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h or Emu(9525))
        s.fill.solid()
        s.fill.fore_color.rgb = self.rgb(color)
        s.line.fill.background()
        return s

    def page(self, i, title, sub=""):
        from pptx.util import Emu, Pt
        sl = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = self.rgb(PAPER)
        M = Emu(609600)
        self.text(sl, M, Emu(330000), Emu(1600000), Emu(300000), f"{['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII'][i]} · {LATIN[i].upper()}", 10, GOLD, font=SERIF)
        self.text(sl, M, Emu(560000), self.W - 2 * M, Emu(620000), title, 28, INK, font=SERIF)
        if sub:
            self.text(sl, M, Emu(1130000), self.W - 2 * M, Emu(300000), sub, 11, MUTE)
        self.rule(sl, M, Emu(1460000), self.W - 2 * M)
        self.text(sl, M, self.H - Emu(430000), Emu(6000000), Emu(260000), "Stoa · 仅供参考，不构成投资建议", 8, MUTE)
        self.text(sl, self.W - M - Emu(900000), self.H - Emu(430000), Emu(900000), Emu(260000), f"{i + 1} / 8", 8, MUTE, font=SERIF)
        sl.name = title
        _ = Pt
        return sl

    def table(self, sl, x, y, w, cols, rows, size=11, max_rows=10, max_cols=8, max_chars=38):
        from pptx.util import Emu, Pt
        cols = [str(c) for c in cols][:max_cols]
        rows = [list(r)[:len(cols)] + [""] * (len(cols) - len(list(r)[:len(cols)])) for r in rows][:max_rows]
        rh = Emu(int(Pt(size).emu * 2.3))
        t = sl.shapes.add_table(len(rows) + 1, len(cols), x, y, w, rh * (len(rows) + 1)).table
        for j, c in enumerate(cols):
            cell = t.cell(0, j)
            cell.text = c
            self._cell(cell, size, True, GOLD)
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r):
                s = _num_text(v) if not isinstance(v, str) else v
                cell = t.cell(i, j)
                cell.text = s if len(s) <= max_chars else s[:max_chars - 1] + "…"
                self._cell(cell, size, False, INK)
        # 列宽按内容长度分配（中文按 2 个字符宽），文字列宽、数字列窄
        def wlen(x):
            x = str(x)
            return sum(2 if ord(ch) > 0x2E80 else 1 for ch in x)
        weights = [min(36, max(6, max([wlen(c)] + [wlen(t.cell(i, j).text) for i in range(1, len(rows) + 1)]))) for j, c in enumerate(cols)]
        tot = sum(weights)
        for j, wt in enumerate(weights):
            t.columns[j].width = Emu(int(w * wt / tot))
        return t

    def _cell(self, cell, size, bold, color):
        from pptx.util import Pt
        cell.fill.solid()
        cell.fill.fore_color.rgb = self.rgb("EFECE6" if bold else "FBF9F4")
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.name = Pt(size), bold, SANS
                r.font.color.rgb = self.rgb(color)

    def empty(self, sl, why=EMPTY):
        from pptx.util import Emu
        self.text(sl, Emu(609600), Emu(2600000), self.W - Emu(1219200), Emu(600000), why, 16, MUTE, italic=True)


def build(deck: dict, path: Path) -> dict:
    """写 pptx；返回 {pages: [...], empty: [...], filled: [...]}。"""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Pt

    d = autofill(deck)
    D = _Deck()
    M = Emu(609600)
    CW = D.W - 2 * M
    top = Emu(1700000)
    name = str(d.get("name") or d.get("symbol") or d.get("title") or "")
    sym = str(d.get("symbol") or "")
    date = str(d.get("date") or "")
    empty = []

    # 1 结论
    c = d.get("conclusion") or {}
    sl = D.page(0, str(d.get("title") or f"{name}{f'（{sym}）' if sym and sym != name else ''} 投研简报"),
                " · ".join(x for x in (date and f"截至 {date}", c.get("source")) if x))
    if c.get("view"):
        D.text(sl, M, top, Emu(3400000), Emu(1200000), str(c["view"]), 60, VIEW_COLOR.get(str(c["view"]), INK), font=SERIF)
        conf = c.get("confidence")
        D.text(sl, M, top + Emu(1250000), Emu(3400000), Emu(400000),
               [x for x in (f"置信度 {conf}%" if conf is not None else "", f"约 {c['horizon_days']} 个交易日" if c.get("horizon_days") else "") if x] or [""], 14, MUTE)
        D.text(sl, M, top + Emu(2000000), Emu(3400000), Emu(800000), "择时判断经回测不具统计显著性，方向是次要信息", 10, MUTE, italic=True)
        rs = [str(x) for x in (c.get("reasons") or [])][:3]
        D.text(sl, M + Emu(3900000), top, CW - Emu(3900000), Emu(300000), "三条依据", 12, GOLD, font=SERIF)
        D.text(sl, M + Emu(3900000), top + Emu(400000), CW - Emu(3900000), Emu(3600000), [f"{i}. {r}" for i, r in enumerate(rs, 1)] or ["（未提供依据）"], 15)
    else:
        D.empty(sl)
        empty.append(TITLES[0])

    # 2 价格与估值
    pv = d.get("price") or {}
    sl = D.page(1, "价格与估值", str(pv.get("source") or ""))
    metrics = pv.get("metrics") or []
    hist = pv.get("history") or []
    if metrics or hist:
        tw = CW if not hist else Emu(int(CW * .42))
        for i, m in enumerate(metrics[:8]):
            lab, val = (m[0], m[1]) if isinstance(m, (list, tuple)) else (m.get("label"), m.get("value"))
            note = (m[2] if isinstance(m, (list, tuple)) and len(m) > 2 else m.get("note") if isinstance(m, dict) else "") or ""
            col, row = i % 2, i // 2
            x, y = M + Emu(int(tw / 2 * col)), top + Emu(1050000 * row)
            D.text(sl, x, y, Emu(int(tw / 2) - Emu(100000)), Emu(300000), str(lab), 10, MUTE)
            D.text(sl, x, y + Emu(260000), Emu(int(tw / 2) - Emu(100000)), Emu(500000), _num_text(val), 24, INK, font=SERIF)
            if note:
                D.text(sl, x, y + Emu(720000), Emu(int(tw / 2) - Emu(100000)), Emu(300000), str(note), 9, MUTE)
        if hist:
            cd = CategoryChartData()
            cd.categories = [str(h[0]) for h in hist]
            cd.add_series("收盘价", [float(h[1]) for h in hist])
            gx = M + tw + Emu(200000) if metrics else M
            ch = sl.shapes.add_chart(XL_CHART_TYPE.LINE, gx, top, D.W - M - gx, Emu(4000000), cd).chart
            ch.has_legend = False
            ch.has_title = False
            ch.plots[0].series[0].format.line.color.rgb = D.rgb(INK)
            ch.plots[0].series[0].smooth = False
            ch.category_axis.tick_labels.font.size = Pt(8)
            ch.value_axis.tick_labels.font.size = Pt(8)
    else:
        D.empty(sl)
        empty.append(TITLES[1])

    # 3 财务
    fin = d.get("financials") or {}
    sl = D.page(2, "财务", str(fin.get("source") or fin.get("note") or ""))
    if fin.get("rows"):
        D.table(sl, M, top, CW, fin.get("columns") or [], fin["rows"], 12, max_rows=9)
    else:
        D.empty(sl)
        empty.append(TITLES[2])

    # 4 利润结构瀑布（原生堆积柱：透明底座 + 增/减两组，首尾为合计）
    wf = d.get("waterfall") or {}
    items = wf.get("items") or []
    sl = D.page(3, "利润结构", str(wf.get("source") or wf.get("period") or ""))
    if len(items) >= 2:
        labels, base, up, down, tot = [], [], [], [], []
        run = 0.0
        for k, it in enumerate(items):
            lab, val = (it[0], float(it[1])) if isinstance(it, (list, tuple)) else (it["label"], float(it["value"]))
            is_total = (it.get("total") if isinstance(it, dict) else None) or k == 0 or (k == len(items) - 1 and wf.get("last_total", True))
            labels.append(str(lab))
            if is_total:
                base.append(0.0), up.append(0.0), down.append(0.0), tot.append(val)
                run = val
            else:
                nxt = run + val
                base.append(min(run, nxt)), tot.append(0.0)
                (up if val >= 0 else down).append(abs(val))
                (down if val >= 0 else up).append(0.0)
                run = nxt
        unit = float(wf.get("unit") or (1e8 if max(abs(x) for x in tot + up + down) >= 1e8 else 1))
        cd = CategoryChartData()
        cd.categories = labels
        for nm, vals in (("底座", base), ("合计", tot), ("增加", up), ("减少", down)):
            cd.add_series(nm, [round(v / unit, 4) for v in vals])
        ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, M, top, CW, Emu(4200000), cd).chart
        ch.has_legend = False
        ch.has_title = False
        ch.plots[0].gap_width = 60
        for s, col in zip(ch.plots[0].series, (None, INK, TERRA, OLIVE)):
            if col is None:
                s.format.fill.background()
                continue
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = D.rgb(col)
            dl = s.data_labels  # 柱上标数值（零值不显示）
            dl.number_format, dl.number_format_is_linked = '#,##0.00;-#,##0.00;;', False
            dl.show_value = True
            dl.font.size = Pt(9)
            dl.font.color.rgb = D.rgb(PAPER if col == INK else INK)
        ch.category_axis.tick_labels.font.size = Pt(9)
        ch.value_axis.tick_labels.font.size = Pt(8)
        D.text(sl, M, top + Emu(4250000), CW, Emu(300000), f"单位：{'亿元' if unit == 1e8 else wf.get('unit_name') or '元'}；墨色为合计，赤陶为增加、橄榄为减少（相对上一步）", 9, MUTE)
    else:
        D.empty(sl)
        empty.append(TITLES[3])

    # 5 Agora 辩论摘要
    db = d.get("debate") or {}
    sl = D.page(4, "多空辩论摘要", str(db.get("source") or ""))
    if db.get("bull") or db.get("bear") or db.get("verdict"):
        half = Emu(int(CW / 2) - Emu(150000))
        D.text(sl, M, top, half, Emu(300000), "多方 · Bull", 13, TERRA, font=SERIF)
        D.text(sl, M, top + Emu(380000), half, Emu(2600000), [f"· {x}" for x in db.get("bull") or []] or ["—"], 13)
        D.text(sl, M + half + Emu(300000), top, half, Emu(300000), "空方 · Bear", 13, OLIVE, font=SERIF)
        D.text(sl, M + half + Emu(300000), top + Emu(380000), half, Emu(2600000), [f"· {x}" for x in db.get("bear") or []] or ["—"], 13)
        D.rule(sl, M, top + Emu(3050000), CW, LINE)
        D.text(sl, M, top + Emu(3150000), CW, Emu(900000), [x for x in (f"裁决：{db['verdict']}" if db.get("verdict") else "", db.get("risk") or "") if x], 12)
    else:
        D.empty(sl)
        empty.append(TITLES[4])

    # 6 风险清单
    rk = d.get("risks") or []
    sl = D.page(5, "风险清单")
    if rk:
        rows = [[(r.get("level") or "") if isinstance(r, dict) else "", (r.get("text") if isinstance(r, dict) else str(r))] for r in rk][:9]
        D.table(sl, M, top, CW, ["级别", "风险"], rows, 12, max_chars=70)
        t = sl.shapes[-1].table
        t.columns[0].width = Emu(1400000)
        t.columns[1].width = CW - Emu(1400000)
    else:
        D.empty(sl)
        empty.append(TITLES[5])

    # 7 同行对比（sweep_filings 矩阵）
    pe = d.get("peers") or {}
    sl = D.page(6, "同行对比", str(pe.get("source") or ""))
    if pe.get("rows"):
        cols = [c for c in pe.get("columns") or [] if c not in ("页码链接",)]
        idx = [i for i, c in enumerate(pe.get("columns") or []) if c in cols]
        rows = [[r[i] if i < len(r) else "" for i in idx] for r in pe["rows"]]
        D.table(sl, M, top, CW, cols, rows, 10, max_rows=10, max_cols=8, max_chars=26)
        cov = pe.get("coverage") or {}
        if cov:
            D.text(sl, M, D.H - Emu(800000), CW, Emu(300000),
                   f"覆盖：请求 {cov.get('requested', '—')} 份，取到 {cov.get('fetched', '—')} 份" + (f"；失败 {'、'.join(map(str, cov.get('failed') or []))}" if cov.get("failed") else ""), 9, MUTE)
    else:
        D.empty(sl)
        empty.append(TITLES[6])

    # 8 失效条件与复盘日期
    iv = d.get("invalidation") or {}
    sl = D.page(7, "失效条件与复盘日期", str(iv.get("source") or ""))
    if iv.get("conditions") or iv.get("review_date"):
        D.text(sl, M, top, Emu(3400000), Emu(300000), "复盘日期", 12, GOLD, font=SERIF)
        D.text(sl, M, top + Emu(350000), Emu(3600000), Emu(800000), str(iv.get("review_date") or "未定"), 36, INK, font=SERIF)
        if iv.get("call"):
            D.text(sl, M, top + Emu(1250000), Emu(3400000), Emu(600000), str(iv["call"]), 10, MUTE)
        D.text(sl, M + Emu(3900000), top, CW - Emu(3900000), Emu(300000), "失效条件 · 出现任一条即重新评估", 12, GOLD, font=SERIF)
        D.text(sl, M + Emu(3900000), top + Emu(400000), CW - Emu(3900000), Emu(3000000), [f"· {x}" for x in iv.get("conditions") or []] or ["—"], 14)
    else:
        D.empty(sl)
        empty.append(TITLES[7])
    srcs = [str(x) for x in d.get("sources") or []][:6]
    if srcs:
        D.text(sl, M, D.H - Emu(1100000), CW, Emu(600000), ["出处："] + srcs, 8, MUTE)
    if d.get("_filled"):
        D.text(sl, M, D.H - Emu(700000), CW, Emu(260000), "由本机记录补入：" + "、".join(d["_filled"]), 8, MUTE)
    path.parent.mkdir(parents=True, exist_ok=True)
    D.prs.save(path)
    warns = []
    if len(items) >= 3:  # 瀑布首尾勾稽：首项 + 中间增减 ≈ 末项，否则提示（不改数）
        vals = [float(it[1] if isinstance(it, (list, tuple)) else it["value"]) for it in items]
        gap = vals[0] + sum(vals[1:-1]) - vals[-1]
        if abs(gap) > max(abs(vals[-1]) * .01, 1e-9):
            warns.append(f"利润瀑布首项加中间增减与末项差 {_num_text(gap)}（可能漏了科目，如少数股东损益、其他收益）")
    return {"pages": len(D.prs.slides), "empty": empty, "filled": d.get("_filled") or [], "deck": d, "warnings": warns}


# ---------------------------------------------------------------- 读回校验
def _deck_numbers(d: dict) -> list[str]:
    """输入里应当出现在页面上的数字（按页面的显示格式），用于读回比对。"""
    out = []
    pv = d.get("price") or {}
    for m in pv.get("metrics") or []:
        v = m[1] if isinstance(m, (list, tuple)) else m.get("value")
        out.append(_num_text(v))
    for r in (d.get("financials") or {}).get("rows") or []:
        out += [_num_text(v) if not isinstance(v, str) else v[:37] for v in r[1:8]]
    c = d.get("conclusion") or {}
    if c.get("confidence") is not None:
        out.append(f"{c['confidence']}%")
    iv = d.get("invalidation") or {}
    if iv.get("review_date"):
        out.append(str(iv["review_date"]))
    return [x for x in out if x and x != "—"]


def verify(path: Path, deck: dict) -> dict:
    from pptx import Presentation
    prs = Presentation(str(path))
    texts, chart_vals = [], []
    titles = []
    for sl in prs.slides:
        t = []
        for sh in sl.shapes:
            if sh.has_text_frame:
                t.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                t += [c.text for row in sh.table.rows for c in row.cells]
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for s in sh.chart.plots[0].series:
                    chart_vals += [v for v in s.values if v is not None]
        texts.append("\n".join(t))
        titles.append(sl.name)
    full = "\n".join(texts)
    miss = [x for x in _deck_numbers(deck) if x not in full]
    wf = (deck.get("waterfall") or {}).get("items") or []
    wf_ok = True
    if len(wf) >= 2:
        vals = [float(it[1] if isinstance(it, (list, tuple)) else it["value"]) for it in wf]
        unit = float((deck.get("waterfall") or {}).get("unit") or (1e8 if max(abs(v) for v in vals) >= 1e8 else 1))
        for v in vals:
            if not any(abs(abs(v) / unit - cv) < 1e-3 * max(1, abs(v) / unit) for cv in chart_vals):
                wf_ok = False
                miss.append(f"瀑布 {v}")
    hist = (deck.get("price") or {}).get("history") or []
    for h in hist[:3] + hist[-3:]:
        if not any(abs(float(h[1]) - cv) < 1e-6 for cv in chart_vals):
            miss.append(f"走势 {h}")
    ok = len(prs.slides) == 8 and titles[1:] == TITLES[1:] and not miss
    return {"ok": ok, "slides": len(prs.slides), "titles": titles, "missing": miss, "waterfall_ok": wf_ok}
